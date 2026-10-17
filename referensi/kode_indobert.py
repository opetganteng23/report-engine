"""Metode 5: fine-tuning IndoBERT (indobenchmark/indobert-base-p1) untuk kategori, sentimen, dan keluhan.

Data latih : dataset final 10.000 tweet berlabel LLM (hasil_v3_10000_final.jsonl), hanya baris Metode = LLM,
             DIKURANGI seluruh sampel uji (385 lama + 615 tambahan) agar tidak bocor ke evaluasi.
Validasi   : 10% data latih (berstrata, benih tetap) -> loss & akurasi per epoch (grafik epoch).
Model      : model dari epoch dengan macro-F1 validasi tertinggi disimpan di model/<tugas>/.

Pemakaian (dari folder crawling-lagi):
  python metode_5_indobert/indobert.py latih            -> latih 3 model, simpan riwayat epoch + grafik
  python metode_5_indobert/indobert.py prediksi         -> tulis kolom "<Tugas> IndoBERT" ke file label engine uji
"""

import json
import logging
import random
import sys
import time
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))  # agar bisa dijalankan langsung

import numpy as np
import torch
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font
from sklearn.metrics import accuracy_score, f1_score
from sklearn.model_selection import train_test_split
from torch.utils.data import DataLoader
from transformers import AutoModelForSequenceClassification, AutoTokenizer, get_linear_schedule_with_warmup

logging.basicConfig(level=logging.INFO, format="%(message)s")
for nama in ("transformers", "httpx", "httpx2"):
    logging.getLogger(nama).setLevel(logging.ERROR)
logger = logging.getLogger("indobert")

FOLDER = Path(__file__).resolve().parent
KERJA = Path.home() / "Downloads" / "Hasil_Penelitian" / "_file_kerja"
FILE_LATIH = KERJA / "analisis_llm" / "hasil_v3_10000_final.jsonl"
FILE_UJI_LAMA = KERJA / "evaluasi_akurasi" / "lembar_label_v2_pelabel1.xlsx"
FILE_UJI_TAMBAHAN = KERJA / "evaluasi_akurasi" / "kunci_tambahan615.json"
FOLDER_HASIL = KERJA / "evaluasi_akurasi" / "indobert"

NAMA_MODEL = "indobenchmark/indobert-base-p1"
LABEL = {
    "Kategori": ["Angkutan Umum", "Lahan Parkir", "Manajemen Lalu Lintas", "Tidak Relevan"],
    "Sentimen": ["Positif", "Negatif", "Netral"],
    "Keluhan": ["Ya", "Tidak"],
}
PARAM = {"epoch": 4, "batch": 32, "lr": 2e-5, "max_len": 128, "warmup": 0.1, "weight_decay": 0.01,
         "porsi_validasi": 0.1, "benih": 2026}
TEBAL = Font(bold=True)


def atur_benih(benih: int) -> None:
    """Benih acak tetap untuk Python, NumPy, dan PyTorch (hasil bisa diulang)."""
    random.seed(benih)
    np.random.seed(benih)
    torch.manual_seed(benih)
    torch.cuda.manual_seed_all(benih)


def teks_uji() -> set[str]:
    """Teks seluruh sampel uji (385 lama + 615 tambahan) yang tidak boleh ikut dilatih."""
    lama = {b[2] for b in list(load_workbook(FILE_UJI_LAMA, read_only=True)["Pelabelan"].iter_rows(values_only=True))[1:]}
    tambahan = {k["teks_bersih"] for k in json.loads(FILE_UJI_TAMBAHAN.read_text(encoding="utf-8"))}
    return lama | tambahan


def muat_data_latih() -> tuple[list[str], dict[str, list[str]], int]:
    """Teks + label LLM per tugas, tanpa sampel uji."""
    uji = teks_uji()
    teks, label, keluar = [], {t: [] for t in LABEL}, 0
    with FILE_LATIH.open(encoding="utf-8") as f:
        for d in map(json.loads, f):
            if d["Metode"] != "LLM" or not d["teks_bersih"].strip():
                continue
            if d["teks_bersih"] in uji:
                keluar += 1
                continue
            teks.append(d["teks_bersih"])
            for t in LABEL:
                label[t].append(d[t])
    return teks, label, keluar


def buat_loader(tokenizer, teks: list[str], y: list[int] | None, acak: bool) -> DataLoader:
    """Tokenisasi lalu bungkus jadi DataLoader."""
    enc = tokenizer(teks, truncation=True, max_length=PARAM["max_len"], padding="max_length", return_tensors="pt")
    tensor = [enc["input_ids"], enc["attention_mask"]] + ([torch.tensor(y)] if y is not None else [])
    data = torch.utils.data.TensorDataset(*tensor)
    return DataLoader(data, batch_size=PARAM["batch"], shuffle=acak)


@torch.no_grad()
def tebak(model, loader, perangkat) -> tuple[list[int], float]:
    """Prediksi kelas (+ rata-rata loss jika label tersedia)."""
    model.eval()
    hasil, total_loss, n = [], 0.0, 0
    for batch in loader:
        ids, mask = batch[0].to(perangkat), batch[1].to(perangkat)
        with torch.autocast(device_type=perangkat.type, dtype=torch.float16, enabled=perangkat.type == "cuda"):
            keluaran = model(input_ids=ids, attention_mask=mask,
                             labels=batch[2].to(perangkat) if len(batch) > 2 else None)
        if len(batch) > 2:
            total_loss += keluaran.loss.item() * len(ids)
            n += len(ids)
        hasil.extend(keluaran.logits.argmax(-1).tolist())
    return hasil, (total_loss / n if n else 0.0)


def latih_satu(tugas: str, teks: list[str], label: list[str], tokenizer, perangkat) -> list[dict]:
    """Fine-tune satu model; kembalikan riwayat per epoch; simpan model epoch terbaik."""
    kelas = LABEL[tugas]
    y = [kelas.index(v) for v in label]
    t_lat, t_val, y_lat, y_val = train_test_split(teks, y, test_size=PARAM["porsi_validasi"], stratify=y,
                                                  random_state=PARAM["benih"])
    atur_benih(PARAM["benih"])
    model = AutoModelForSequenceClassification.from_pretrained(NAMA_MODEL, num_labels=len(kelas)).to(perangkat)
    dl_lat = buat_loader(tokenizer, t_lat, y_lat, acak=True)
    dl_val = buat_loader(tokenizer, t_val, y_val, acak=False)
    opt = torch.optim.AdamW(model.parameters(), lr=PARAM["lr"], weight_decay=PARAM["weight_decay"])
    langkah = len(dl_lat) * PARAM["epoch"]
    jadwal = get_linear_schedule_with_warmup(opt, int(langkah * PARAM["warmup"]), langkah)
    skala = torch.amp.GradScaler(enabled=perangkat.type == "cuda")
    riwayat, terbaik = [], -1.0
    for epoch in range(1, PARAM["epoch"] + 1):
        model.train()
        mulai, total, n = time.time(), 0.0, 0
        for ids, mask, yb in dl_lat:
            ids, mask, yb = ids.to(perangkat), mask.to(perangkat), yb.to(perangkat)
            opt.zero_grad()
            with torch.autocast(device_type=perangkat.type, dtype=torch.float16, enabled=perangkat.type == "cuda"):
                loss = model(input_ids=ids, attention_mask=mask, labels=yb).loss
            skala.scale(loss).backward()
            skala.unscale_(opt)
            torch.nn.utils.clip_grad_norm_(model.parameters(), 1.0)
            skala.step(opt)
            skala.update()
            jadwal.step()
            total += loss.item() * len(ids)
            n += len(ids)
        pred_lat, _ = tebak(model, buat_loader(tokenizer, t_lat[:2000], y_lat[:2000], acak=False), perangkat)
        pred_val, loss_val = tebak(model, dl_val, perangkat)
        catat = {
            "tugas": tugas, "epoch": epoch, "loss_latih": total / n, "loss_validasi": loss_val,
            "akurasi_latih": accuracy_score(y_lat[:2000], pred_lat), "akurasi_validasi": accuracy_score(y_val, pred_val),
            "macro_f1_validasi": f1_score(y_val, pred_val, average="macro"), "detik": time.time() - mulai,
        }
        riwayat.append(catat)
        logger.info("  %s epoch %d: loss latih %.4f | loss val %.4f | akurasi val %.1f%% | macro-F1 val %.1f%% | %.0f dtk",
                    tugas, epoch, catat["loss_latih"], loss_val, catat["akurasi_validasi"] * 100,
                    catat["macro_f1_validasi"] * 100, catat["detik"])
        if catat["macro_f1_validasi"] > terbaik:
            terbaik = catat["macro_f1_validasi"]
            model.save_pretrained(FOLDER / "model" / tugas.lower())
            tokenizer.save_pretrained(FOLDER / "model" / tugas.lower())
    for c in riwayat:
        c["terbaik"] = c["macro_f1_validasi"] == terbaik
    del model
    torch.cuda.empty_cache()
    return riwayat


def grafik_epoch(riwayat: list[dict]) -> Path:
    """Grafik loss & akurasi per epoch untuk ketiga tugas (PNG 300 dpi)."""
    import matplotlib
    matplotlib.use("Agg")
    import matplotlib.pyplot as plt
    fig, axs = plt.subplots(2, 3, figsize=(13, 7), sharex=True)
    for j, tugas in enumerate(LABEL):
        r = [c for c in riwayat if c["tugas"] == tugas]
        ep = [c["epoch"] for c in r]
        for ax, (kunci_a, kunci_b, judul) in zip(axs[:, j], [("loss_latih", "loss_validasi", "Loss"),
                                                              ("akurasi_latih", "akurasi_validasi", "Akurasi")]):
            ax.plot(ep, [c[kunci_a] for c in r], "o-", color="#2a78d6", lw=2, ms=6, label="Latih")
            ax.plot(ep, [c[kunci_b] for c in r], "s-", color="#eb6834", lw=2, ms=6, label="Validasi")
            terbaik = next(c for c in r if c["terbaik"])
            ax.axvline(terbaik["epoch"], color="#9aa3af", lw=1, ls="--")
            ax.set_title(f"{tugas} – {judul}", loc="left", fontsize=10, fontweight="bold")
            ax.set_xticks(ep)
            ax.grid(axis="y", color="#e3e5e8", lw=0.6)
            for sisi in ("top", "right"):
                ax.spines[sisi].set_visible(False)
        axs[1, j].set_xlabel("Epoch")
    axs[0, 0].legend(frameon=False)
    fig.suptitle("IndoBERT: loss dan akurasi per epoch (garis putus = epoch terbaik menurut macro-F1 validasi)",
                 x=0.01, ha="left", fontweight="bold")
    fig.tight_layout()
    FOLDER_HASIL.mkdir(parents=True, exist_ok=True)
    tujuan = FOLDER_HASIL / "8_epoch_indobert.png"
    fig.savefig(tujuan, dpi=300, bbox_inches="tight")
    plt.close(fig)
    return tujuan


def simpan_riwayat(riwayat: list[dict], n_latih: int, keluar: int) -> Path:
    """Excel riwayat pelatihan per epoch + konfigurasi."""
    buku = Workbook()
    sheet = buku.active
    sheet.title = "Riwayat Epoch"
    kolom = ["tugas", "epoch", "loss_latih", "loss_validasi", "akurasi_latih", "akurasi_validasi",
             "macro_f1_validasi", "detik", "terbaik"]
    sheet.append(kolom)
    for c in riwayat:
        sheet.append([round(c[k], 4) if isinstance(c[k], float) else c[k] for k in kolom])
    for sel in sheet[1]:
        sel.font = TEBAL
    konf = buku.create_sheet("Konfigurasi")
    for b in [["Model", NAMA_MODEL], ["Data latih", f"{n_latih} tweet berlabel LLM ({keluar} sampel uji dikeluarkan)"],
              ["Perangkat", torch.cuda.get_device_name(0) if torch.cuda.is_available() else "CPU"]] + \
             [[k, str(v)] for k, v in PARAM.items()]:
        konf.append(b)
    konf.column_dimensions["A"].width = 18
    konf.column_dimensions["B"].width = 70
    FOLDER_HASIL.mkdir(parents=True, exist_ok=True)
    tujuan = FOLDER_HASIL / "riwayat_pelatihan_indobert.xlsx"
    buku.save(tujuan)
    return tujuan


def latih() -> None:
    """Latih ketiga model, simpan riwayat + grafik epoch."""
    perangkat = torch.device("cuda" if torch.cuda.is_available() else "cpu")
    logger.info("Perangkat: %s", torch.cuda.get_device_name(0) if perangkat.type == "cuda" else "CPU")
    teks, label, keluar = muat_data_latih()
    logger.info("Data latih: %d tweet (sampel uji dikeluarkan: %d)", len(teks), keluar)
    tokenizer = AutoTokenizer.from_pretrained(NAMA_MODEL)
    riwayat = []
    for tugas in LABEL:
        logger.info("== Melatih model %s", tugas)
        riwayat += latih_satu(tugas, teks, label[tugas], tokenizer, perangkat)
    logger.info("Riwayat: %s | Grafik: %s", simpan_riwayat(riwayat, len(teks), keluar), grafik_epoch(riwayat))


def prediksi(file_uji: Path) -> None:
    """Prediksi sampel uji (sheet 'Label Engine', kolom Teks) -> kolom '<Tugas> IndoBERT'."""
    perangkat = torch.device("cuda" if torch.cuda.is_available() else "cpu")
    buku = load_workbook(file_uji)
    sheet = buku["Label Engine"]
    teks = [r[1] or "" for r in sheet.iter_rows(min_row=2, values_only=True)]
    judul = [c.value for c in sheet[1]]
    for tugas in LABEL:
        folder = FOLDER / "model" / tugas.lower()
        tokenizer = AutoTokenizer.from_pretrained(folder)
        model = AutoModelForSequenceClassification.from_pretrained(folder).to(perangkat)
        pred, _ = tebak(model, buat_loader(tokenizer, teks, None, acak=False), perangkat)
        nama_kolom = f"{tugas} IndoBERT"
        kolom = judul.index(nama_kolom) + 1 if nama_kolom in judul else len(judul) + 1
        if nama_kolom not in judul:
            judul.append(nama_kolom)
        sheet.cell(row=1, column=kolom, value=nama_kolom).font = TEBAL
        for r, p in enumerate(pred, 2):
            sheet.cell(row=r, column=kolom, value=LABEL[tugas][p])
        del model
        torch.cuda.empty_cache()
    buku.save(file_uji)
    logger.info("Prediksi IndoBERT ditulis ke %s", file_uji)


if __name__ == "__main__":
    perintah = sys.argv[1] if len(sys.argv) > 1 else ""
    if perintah == "latih":
        latih()
    elif perintah == "prediksi":
        prediksi(Path(sys.argv[2]) if len(sys.argv) > 2 else KERJA / "evaluasi_akurasi" / "hasil_engine_sampel_v3.xlsx")
    else:
        print(__doc__)
