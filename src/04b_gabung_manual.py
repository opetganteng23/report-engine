"""Tahap 4b: Gabung label manual menjadi data uji gold.

Membaca lembar manual yang sudah diisi, memvalidasi label, menghitung
Cohen's kappa per tugas (bila lebih dari satu anotator), lalu menulis
data/split/test_gold.jsonl dan jumlah label manual per kelas.

Jika ada lebih dari satu anotator dan labelnya berbeda, skrip berhenti dan
menulis daftar ketidaksepakatan; cara penentuan label gold belum ditetapkan.

Pemakaian: python src/04b_gabung_manual.py --config config.yaml
"""

import argparse
import importlib.util
import itertools
import json
import sys
from pathlib import Path

import yaml
from openpyxl import load_workbook
from sklearn.metrics import cohen_kappa_score

spec = importlib.util.spec_from_file_location("split", Path(__file__).with_name("04_split.py"))
split = importlib.util.module_from_spec(spec)
spec.loader.exec_module(split)

TUGAS = {"Kategori": "kategori", "Sentimen": "sentimen", "Keluhan": "keluhan"}


def read_sheet(path, labels):
    ws = load_workbook(path, read_only=True).active
    rows = ws.iter_rows(values_only=True)
    header = list(next(rows))
    if header[:len(split.KOLOM_MANUAL)] != split.KOLOM_MANUAL:
        sys.exit(f"Header {path.name} tidak sesuai: {header}")
    hasil, salah = {}, []
    for row in rows:
        r = dict(zip(header, row))
        if r["ID Sampel"] is None:
            continue
        lab = {}
        for kol, key in TUGAS.items():
            v = (str(r[kol]).strip() if r[kol] is not None else None)
            if v not in labels[key]:
                salah.append(f"{path.name} {r['ID Sampel']} {kol}={r[kol]!r}")
            lab[kol] = v
        lab["Catatan"] = r["Catatan"]
        hasil[r["ID Sampel"]] = lab
    if salah:
        sys.exit("Label manual kosong atau tidak valid:\n" + "\n".join(salah))
    return hasil


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--config", required=True)
    args = ap.parse_args()

    cfg_path = Path(args.config).resolve()
    engine_dir = cfg_path.parent
    cfg = yaml.safe_load(cfg_path.read_text(encoding="utf-8"))
    sp = cfg["split"]
    labels = cfg["labels"]

    with open(engine_dir / sp["test"], encoding="utf-8") as f:
        test = [json.loads(l) for l in f if l.strip()]
    ids = [r["id_sampel"] for r in test]

    paths = split.manual_paths(engine_dir / sp["test_manual"], sp["n_anotator"])
    sheets = [read_sheet(p, labels) for p in paths]
    for p, s in zip(paths, sheets):
        if set(s) != set(ids):
            sys.exit(f"{p.name}: ID Sampel tidak cocok dengan test.jsonl")

    ringkasan = {"n_anotator": len(sheets), "jumlah_post_uji": len(test)}

    if len(sheets) > 1:
        kappa = {}
        for (i, a), (j, b) in itertools.combinations(enumerate(sheets, 1), 2):
            kappa[f"anotator{i}_vs_anotator{j}"] = {}
            for kol in TUGAS:
                k = cohen_kappa_score([a[x][kol] for x in ids], [b[x][kol] for x in ids])
                # kappa tak terdefinisi (mis. kedua anotator memakai satu kelas saja) -> null
                kappa[f"anotator{i}_vs_anotator{j}"][kol] = None if k != k else round(k, 4)
        ringkasan["cohen_kappa"] = kappa
        beda = [{"id_sampel": x, "tugas": kol, "label": [s[x][kol] for s in sheets]}
                for x in ids for kol in TUGAS if len({s[x][kol] for s in sheets}) > 1]
        ringkasan["jumlah_ketidaksepakatan"] = len(beda)
        if beda:
            out = engine_dir / cfg["paths"]["eval"] / "ketidaksepakatan_tahap4b.json"
            out.write_text(json.dumps({"ringkasan": ringkasan, "detail": beda},
                                      ensure_ascii=False, indent=2), encoding="utf-8")
            print(json.dumps(ringkasan, ensure_ascii=False, indent=2))
            sys.exit(f"Ada {len(beda)} label yang berbeda antar-anotator (lihat {out.name}). "
                     "Aturan penentuan label gold belum ditetapkan.")

    gold_src = sheets[0]
    gold = []
    for r in test:
        g = gold_src[r["id_sampel"]]
        row = {k: v for k, v in r.items() if k not in ("Kategori", "Sentimen", "Keluhan", "metode")}
        row.update({f"{k}_LLM": r[k] for k in TUGAS}, metode_LLM=r["metode"])
        row.update({k: g[k] for k in TUGAS}, catatan_manual=g["Catatan"])
        gold.append(row)

    with open(engine_dir / sp["test_gold"], "w", encoding="utf-8") as f:
        for r in gold:
            f.write(json.dumps(r, ensure_ascii=False) + "\n")

    ringkasan["label_manual_per_kelas"] = {
        kol: {lab: sum(r[kol] == lab for r in gold) for lab in labels[key]} for kol, key in TUGAS.items()}
    (engine_dir / cfg["paths"]["eval"] / "ringkasan_tahap4b.json").write_text(
        json.dumps(ringkasan, ensure_ascii=False, indent=2), encoding="utf-8")
    print(json.dumps(ringkasan, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
